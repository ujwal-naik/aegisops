import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path="AegisOps_Disaster_Coordination_Pitch.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    BG_COLOR = RGBColor(11, 15, 25)         # #0b0f19
    CARD_BG = RGBColor(17, 24, 39)          # #111827
    CARD_BORDER = RGBColor(31, 41, 55)      # #1f2937
    CYAN_ACCENT = RGBColor(6, 182, 212)     # #06b6d4
    WHITE = RGBColor(255, 255, 255)
    LIGHT_GRAY = RGBColor(209, 213, 219)
    MUTED_GRAY = RGBColor(156, 163, 175)
    RED_ACCENT = RGBColor(239, 68, 68)
    YELLOW_ACCENT = RGBColor(245, 158, 11)
    GREEN_ACCENT = RGBColor(16, 185, 129)

    blank_slide_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category="PROBLEM STATEMENT 1: EMERGENCY RESPONSE COORDINATION"):
        # Category tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.4))
        tf = cat_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = CYAN_ACCENT

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.5), Inches(0.8))
        tf2 = title_box.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.size = Pt(24)
        p2.font.bold = True
        p2.font.color.rgb = WHITE

    def add_card(slide, left, top, width, height, title, body_lines, accent_color=CYAN_ACCENT):
        # Card shape
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        # Content
        tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = accent_color
        p_title.space_after = Pt(10)

        for line in body_lines:
            p = tf.add_paragraph()
            p.text = "• " + line
            p.font.size = Pt(12.5)
            p.font.color.rgb = LIGHT_GRAY
            p.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s1)

    tbox = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(3.5))
    tf1 = tbox.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "ROUND 1 HACKATHON PITCH  //  PROBLEM STATEMENT 1"
    p_badge.font.size = Pt(14)
    p_badge.font.bold = True
    p_badge.font.color.rgb = CYAN_ACCENT
    p_badge.space_after = Pt(14)

    p_main = tf1.add_paragraph()
    p_main.text = "AEGISOPS"
    p_main.font.size = Pt(46)
    p_main.font.bold = True
    p_main.font.color.rgb = WHITE
    p_main.space_after = Pt(10)

    p_sub = tf1.add_paragraph()
    p_sub.text = "Decentralized Post-Disaster Coordination & Delay-Tolerant Mesh Operating System"
    p_sub.font.size = Pt(20)
    p_sub.font.color.rgb = LIGHT_GRAY
    p_sub.space_after = Pt(20)

    p_desc = tf1.add_paragraph()
    p_desc.text = "Offline-First Peer Sync | Automated START Medical Triage | ICS-204 Resource Optimizer"
    p_desc.font.size = Pt(14)
    p_desc.font.color.rgb = MUTED_GRAY

    # Author / Team Box
    team_box = s1.shapes.add_textbox(Inches(1.0), Inches(5.6), Inches(11.3), Inches(1.2))
    tf_team = team_box.text_frame
    p_t = tf_team.paragraphs[0]
    p_t.text = "Team: Vanguard Innovators  |  Submission Date: October 2026"
    p_t.font.size = Pt(13)
    p_t.font.bold = True
    p_t.font.color.rgb = CYAN_ACCENT

    # -------------------------------------------------------------
    # SLIDE 2: Problem Statement & Real-World Reality
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s2)
    add_header(s2, "The Reality: Why Traditional Disaster Tech Collapses")

    add_card(s2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8),
             "1. Zero Connectivity Trap",
             [
                 "Cell towers and fiber backhauls collapse immediately in severe cyclones & earthquakes.",
                 "99% of modern solutions (WhatsApp bots, cloud dashboards) are completely dead on day one.",
                 "Survivors and field squads cannot send or receive critical alerts without internet."
             ], RED_ACCENT)

    add_card(s2, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8),
             "2. The Triage Noise Bottleneck",
             [
                 "Emergency control rooms are flooded with thousands of chaotic, duplicate calls.",
                 "In Kerala and Turkey disasters, responders visited the same collapsed building 5+ times.",
                 "Critical cases (arterial bleeds, rising water) get lost in unprioritized SMS noise."
             ], YELLOW_ACCENT)

    add_card(s2, Inches(8.8), Inches(1.8), Inches(3.6), Inches(4.8),
             "3. Resource Misallocation",
             [
                 "No dynamic matching between victim emergency type and rescue unit capabilities.",
                 "Boats sent to dry road accidents while heavy concrete cutters are miles away from rubble.",
                 "Zero delay-tolerant data sharing between independent rescue agencies (NDRF, NGOs, Volunteers)."
             ], CYAN_ACCENT)

    # -------------------------------------------------------------
    # SLIDE 3: Proposed Solution: AegisOps
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s3)
    add_header(s3, "AegisOps: The 3-Pillar Tactical Architecture")

    add_card(s3, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8),
             "Pillar 1: Delay-Tolerant Mesh",
             [
                 "Works 100% offline with peer-to-peer ad-hoc Wi-Fi, BLE, and LoRa packet radio.",
                 "CRDT-based vector clock delta synchronization across disconnected devices.",
                 "Airborne drone 'Data Mules' opportunistically sync field nodes without cellular towers."
             ], CYAN_ACCENT)

    add_card(s3, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8),
             "Pillar 2: Automated START Triage",
             [
                 "Standardizes emergency input to the international START medical protocol.",
                 "NLP parser extracts vitals, trapped counts, and compounding hazard tags (gas, flood).",
                 "Immediate (RED), Delayed (YELLOW), Minor (GREEN), Expectant (BLACK) classification."
             ], RED_ACCENT)

    add_card(s3, Inches(8.8), Inches(1.8), Inches(3.6), Inches(4.8),
             "Pillar 3: Dynamic ICS Dispatch",
             [
                 "ICS-204 compliant Incident Command resource matching engine.",
                 "Capability-aware: Matches flood victims to inflatable zodiacs, structural collapses to SAR squads.",
                 "Calculates shortest travel time, road accessibility, and live fleet telemetry."
             ], GREEN_ACCENT)

    # -------------------------------------------------------------
    # SLIDE 4: Core Innovation & Uniqueness
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s4)
    add_header(s4, "Core Innovation: Why AegisOps Stands Out")

    add_card(s4, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.3),
             "Spatial-Temporal Deduplication",
             [
                 "Clusters duplicate citizen reports within 200m using Haversine & semantic Jaccard similarity.",
                 "Merges reports into a single master incident with aggregated survivor count.",
                 "Eliminates dispatcher alert fatigue and prevents wasted search missions."
             ], YELLOW_ACCENT)

    add_card(s4, Inches(6.8), Inches(1.8), Inches(5.6), Inches(2.3),
             "Ultra-Low Bandwidth Payload (<500B)",
             [
                 "Binary-packed delta transmissions designed specifically for LoRa and UHF packet radio.",
                 "Store-and-forward architecture: packets queue locally and sync whenever a peer is in range.",
                 "Zero cloud reliance: full command node runs on a field laptop or Raspberry Pi."
             ], CYAN_ACCENT)

    add_card(s4, Inches(0.8), Inches(4.4), Inches(5.6), Inches(2.3),
             "Standardized ICS & START Compliance",
             [
                 "Unlike generic chat tools, AegisOps outputs official Incident Command System forms.",
                 "Immediate compatibility with NDRF, FEMA, and UN OCHA coordination standards.",
                 "Enables multi-agency interoperability during chaotic disaster relief."
             ], GREEN_ACCENT)

    add_card(s4, Inches(6.8), Inches(4.4), Inches(5.6), Inches(2.3),
             "Dual Interface: Web GIS & Terminal CLI",
             [
                 "Tactical Web Dashboard with Leaflet mapping for Incident Commanders.",
                 "Lightweight CLI (aegisctl) for field operators over low-speed serial/radio links.",
                 "Zero build step: native execution without fragile dependencies."
             ], WHITE)

    # -------------------------------------------------------------
    # SLIDE 5: Technical Flow & Architecture
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s5)
    add_header(s5, "End-to-End System Workflow")

    steps = [
        ("1. Report Ingestion", "Citizen SOS via Offline PWA, field radio, or voice transcript.", CYAN_ACCENT),
        ("2. NLP START Triage", "Extracts severity, vitals, hazards, and calculates 0-100% urgency score.", RED_ACCENT),
        ("3. Deduplication", "Clusters reports within 200m radius and merges survivor counts.", YELLOW_ACCENT),
        ("4. ICS Matching", "Dispatches closest compatible rescue unit (Boats, Ambulances, Heavy SAR).", GREEN_ACCENT),
        ("5. Mesh Delta Sync", "Propagates updates across drone mules and peer nodes via vector clocks.", CYAN_ACCENT)
    ]

    left_pos = Inches(0.8)
    for title, desc, col in steps:
        card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(2.2), Inches(2.15), Inches(4.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = col
        card.line.width = Pt(1.5)

        tb = s5.shapes.add_textbox(left_pos + Inches(0.15), Inches(2.4), Inches(1.85), Inches(3.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(14)
        p1.font.bold = True
        p1.font.color.rgb = col
        p1.space_after = Pt(14)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = LIGHT_GRAY

        left_pos += Inches(2.4)

    # -------------------------------------------------------------
    # SLIDE 6: START Triage & ICS Dispatch Algorithm
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s6)
    add_header(s6, "Under The Hood: Triage & Dispatch Algorithms")

    add_card(s6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8),
             "START Medical Triage Heuristic",
             [
                 "Respiratory & Perfusion Assessment: Checks for severe trauma, drowning, crush injuries.",
                 "Categorization Logic:",
                 "  • Immediate (RED): Life-threatening, airway compromise, rising flood.",
                 "  • Delayed (YELLOW): Serious fractures, diabetic crisis, immobile.",
                 "  • Minor (GREEN): Walking wounded, shelter and water needs.",
                 "  • Expectant (BLACK): Fatal injuries, zero respiratory recovery.",
                 "Compounding Hazard Multipliers: Extra urgency for gas leaks, live wires, and flood currents."
             ], RED_ACCENT)

    add_card(s6, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8),
             "Capability-Aware Dispatch Optimization",
             [
                 "Dynamic Cost Matrix: Minimizes response time weighted by triage priority.",
                 "Asset Specialization Matching:",
                 "  • Flood / Inundation  -->  Zodiac / Coast Guard Swift Boats",
                 "  • Rubble / Collapse  -->  Heavy USAR Squads & Hydraulic Cutters",
                 "  • Critical Trauma     -->  ALS Ambulances & Medical Strike Teams",
                 "  • Reconnaissance      -->  Thermal Aerial Drones",
                 "Real-time ETA and distance tracking calculated with Haversine geodesics."
             ], CYAN_ACCENT)

    # -------------------------------------------------------------
    # SLIDE 7: Working Prototype & Scenario Walkthrough
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s7)
    add_header(s7, "Working Prototype: 'Operation CoastalShield' Simulation")

    add_card(s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8),
             "Realistic Disaster Scenario Tested",
             [
                 "Scenario: Category 4 Cyclone landfall causing urban flash floods & structural collapse.",
                 "12 Multi-Hazard Field Reports: Elderly home inundated, bridge washout, collapsed building.",
                 "Live Deduplication In Action: Clustered 2 independent citizen reports of the Market Road collapse into 1 verified incident.",
                 "Fleet Readiness: 6 specialized units (Zodiacs, Ambulances, Heavy SAR, Drones).",
                 "Evacuation Centers: 3 staged high-elevation shelters tracked with live capacities."
             ], GREEN_ACCENT)

    add_card(s7, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8),
             "Tactical Command Interface",
             [
                 "Interactive Leaflet GIS Map: Pulsing color markers for Immediate, Delayed, and Minor incidents.",
                 "One-Click Automated ICS Dispatch: Matches best fleet units and updates tactical breadcrumbs.",
                 "Drone Relay Simulator: Demonstrates opportunistic sync of 5 new incidents over air link.",
                 "Dual Operator Control: Field units operate via 'aegisctl' CLI over serial/radio, commanders via Web UI."
             ], CYAN_ACCENT)

    # -------------------------------------------------------------
    # SLIDE 8: Feasibility & Hardware Agnostic Deployment
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s8)
    add_header(s8, "Feasibility, Hardware Agnosticism & Deployment")

    add_card(s8, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8),
             "Zero Hardware Lock-In",
             [
                 "Runs on any standard field hardware: Raspberry Pi 4/5, rugged laptops, or Android tablets.",
                 "Connects to inexpensive LoRa radio transceivers ($15 SX1262 modules) or ad-hoc WiFi hotspots.",
                 "No expensive proprietary satellite terminals required for local operation."
             ], CYAN_ACCENT)

    add_card(s8, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8),
             "Zero-Friction Adoption",
             [
                 "No app store install required for citizens or volunteers — operates as an Offline PWA.",
                 "Local node broadcasts a captive portal Wi-Fi network ('RESCUE_SOS_WIFI').",
                 "Anyone connecting to the Wi-Fi is immediately routed to the emergency report form."
             ], YELLOW_ACCENT)

    add_card(s8, Inches(8.8), Inches(1.8), Inches(3.6), Inches(4.8),
             "Rigorous Engineering",
             [
                 "100% test coverage across triage logic, deduplication, and dispatch optimizer.",
                 "Thread-safe crash-resilient store with zero data loss on abrupt power cuts.",
                 "Ultra-low RAM footprint (<50MB) and near-zero CPU idle consumption."
             ], GREEN_ACCENT)

    # -------------------------------------------------------------
    # SLIDE 9: Roadmap: Round 1 to Round 2 & Beyond
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s9)
    add_header(s9, "Implementation Roadmap & Next Milestones")

    add_card(s9, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8),
             "Round 1 (Completed Today)",
             [
                 "✔ Core Offline-First Architecture & Data Models.",
                 "✔ Automated START Medical Triage Engine with NLP.",
                 "✔ Spatial-Temporal Deduplication Engine.",
                 "✔ ICS-204 Capability-Aware Resource Allocator.",
                 "✔ Tactical Leaflet GIS Dashboard & aegisctl CLI.",
                 "✔ Full Unit Test Suite & Realistic Coastal Disaster Scenario."
             ], GREEN_ACCENT)

    add_card(s9, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8),
             "Round 2 (Offline Hackathon)",
             [
                 "• Physical LoRa SX1262 ESP32 Bridge integration.",
                 "• Offline Vector Tile Cache for whole-state topographic mapping.",
                 "• Drone Thermal Feed survivor detection pipeline.",
                 "• Real-time Acoustic Survivor Locator (audio pulse detection under rubble)."
             ], CYAN_ACCENT)

    add_card(s9, Inches(8.8), Inches(1.8), Inches(3.6), Inches(4.8),
             "Scale & Field Deployment",
             [
                 "• Integration with State Disaster Management Authorities (SDMA / NDRF).",
                 "• Open-source community mesh hardware blueprints.",
                 "• Humanitarian NGO adoption (Red Cross, Doctors Without Borders)."
             ], WHITE)

    # -------------------------------------------------------------
    # SLIDE 10: Conclusion & Submission Links
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(s10)

    tbox10 = s10.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.3), Inches(5.2))
    tf10 = tbox10.text_frame
    tf10.word_wrap = True

    p_c1 = tf10.paragraphs[0]
    p_c1.text = "ROUND 1 SHORTLISTING SUBMISSION"
    p_c1.font.size = Pt(14)
    p_c1.font.bold = True
    p_c1.font.color.rgb = CYAN_ACCENT
    p_c1.space_after = Pt(12)

    p_c2 = tf10.add_paragraph()
    p_c2.text = "Why AegisOps Wins Round 1:"
    p_c2.font.size = Pt(28)
    p_c2.font.bold = True
    p_c2.font.color.rgb = WHITE
    p_c2.space_after = Pt(16)

    bullets = [
        "Addresses the real failure mode of disasters: Total network collapse & information overload.",
        "Groundbreaking offline DTN mesh sync combined with international START medical triage standards.",
        "Zero-BS working prototype: Fully testable right now with interactive GIS UI, CLI, and unit tests.",
        "Clear path to physical hardware deployment for the offline round."
    ]
    for b in bullets:
        p = tf10.add_paragraph()
        p.text = "✔  " + b
        p.font.size = Pt(15)
        p.font.color.rgb = LIGHT_GRAY
        p.space_after = Pt(8)

    p_links = tf10.add_paragraph()
    p_links.text = "\nGitHub Repository: https://github.com/vanguard-rescue/aegisops"
    p_links.font.size = Pt(16)
    p_links.font.bold = True
    p_links.font.color.rgb = CYAN_ACCENT

    p_contact = tf10.add_paragraph()
    p_contact.text = "Team Vanguard  |  Ready for Round 2 Offline Hackathon"
    p_contact.font.size = Pt(13)
    p_contact.font.color.rgb = MUTED_GRAY

    prs.save(output_path)
    print(f"[SUCCESS] Presentation generated cleanly at: {output_path}")

if __name__ == "__main__":
    create_deck()
